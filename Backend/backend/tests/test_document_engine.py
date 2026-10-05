from hashlib import sha256
from io import BytesIO

from openpyxl import load_workbook
import pytest
from pydantic import ValidationError

from backend.document_engine import DocumentRenderRequest, compose_blocks, render_document


def request_for(output_format: str, *, text: str = "Hello {{ employee.full_name }}", extra_blocks=None, clauses=None, clause_rules=None):
    blocks = [{"type": "heading", "level": 1, "text": "Offer"}, {"type": "paragraph", "text": text}]
    blocks.extend(extra_blocks or [])
    return DocumentRenderRequest(
        document_code="KR-DOC-TEST",
        title="Test document",
        output_format=output_format,
        design_schema={"page_size": "A4", "footer_text": "KRAVIA PRIVATE LIMITED"},
        content_schema=blocks,
        input_snapshot={"employee": {"full_name": "Ravi <script>alert(1)</script>", "remote": True}},
        clause_snapshot=clauses or {},
        clause_rules=clause_rules or [],
    )


def test_compose_resolves_variables_and_fails_closed_when_missing():
    blocks = compose_blocks(request_for("HTML"))
    assert "Ravi <script>alert(1)</script>" in blocks[1]["text"]

    broken = request_for("HTML", text="{{ employee.missing_value }}")
    with pytest.raises(ValueError, match="variable resolution failed"):
        compose_blocks(broken)


def test_conditional_clause_uses_frozen_clause_snapshot():
    request = request_for(
        "HTML",
        extra_blocks=[{"type": "clause", "clause_code": "REMOTE_WORK"}],
        clauses={"REMOTE_WORK": {"version": 3, "content": "Remote clause for {{ employee.full_name }}", "content_hash": "a" * 64}},
        clause_rules=[{"clause_code": "REMOTE_WORK", "when": {"field": "employee.remote", "truthy": True}}],
    )
    blocks = compose_blocks(request)
    clause = blocks[-1]
    assert clause["clause_code"] == "REMOTE_WORK"
    assert clause["clause_version"] == 3
    assert "Ravi" in clause["text"]

    request.input_snapshot["employee"]["remote"] = False
    blocks = compose_blocks(request)
    assert all(block.get("clause_code") != "REMOTE_WORK" for block in blocks)


def test_all_output_renderers_are_deterministic_binary_outputs():
    signatures = {"PDF": b"%PDF", "DOCX": b"PK", "XLSX": b"PK"}
    for output_format in ("PDF", "DOCX", "HTML", "XLSX"):
        rendered = render_document(request_for(output_format))
        assert len(rendered.content) > 50
        assert rendered.sha256 == sha256(rendered.content).hexdigest()
        assert len(rendered.sha256) == 64
        if output_format in signatures:
            assert rendered.content.startswith(signatures[output_format])


def test_html_escapes_snapshot_values_instead_of_emitting_raw_script():
    rendered = render_document(request_for("HTML"))
    html = rendered.content.decode("utf-8")
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html


def test_xlsx_literals_neutralize_formula_like_content_and_snapshot_values():
    request = DocumentRenderRequest(
        document_code="KR-XLSX-LITERAL",
        title="Spreadsheet literal test",
        output_format="XLSX",
        content_schema=[
            {"type": "heading", "text": "=heading"},
            {"type": "paragraph", "text": "{{ employee.full_name }}"},
            {"type": "key_value", "items": [{"label": "+label", "value": "-value"}]},
            {"type": "table", "rows": [["@cell", "=table"]]},
            {"type": "signature", "name": "=signature"},
        ],
        input_snapshot={"employee": {"full_name": "=snapshot"}},
    )
    sheet = load_workbook(BytesIO(render_document(request).content), data_only=False).active
    values = [cell.value for row in sheet.iter_rows() for cell in row if cell.value]
    assert "'=heading" in values
    assert "'=snapshot" in values
    assert "'+label" in values
    assert "'-value" in values
    assert "'@cell" in values
    assert "'=table" in values
    assert "'=signature" in values
    assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)


def test_document_request_rejects_unbounded_structure_before_rendering():
    deeply_nested: dict[str, object] = {}
    cursor = deeply_nested
    for _ in range(13):
        child: dict[str, object] = {}
        cursor["child"] = child
        cursor = child
    with pytest.raises(ValidationError, match="nesting-depth"):
        DocumentRenderRequest(
            document_code="KR-DEPTH-LIMIT",
            title="Limit",
            output_format="HTML",
            input_snapshot=deeply_nested,
        )

    with pytest.raises(ValidationError, match="text-size"):
        DocumentRenderRequest(
            document_code="KR-TEXT-LIMIT",
            title="Limit",
            output_format="HTML",
            input_snapshot={"note": "x" * 200_001},
        )

    with pytest.raises(ValidationError, match="cell limit"):
        DocumentRenderRequest(
            document_code="KR-CELL-LIMIT",
            title="Limit",
            output_format="XLSX",
            content_schema=[{"type": "table", "rows": [["x"] * 101 for _ in range(100)]}],
        )


def test_document_templates_reject_executable_jinja_expressions():
    with pytest.raises(ValueError, match="dotted variable references only"):
        render_document(request_for("HTML", text="{{ employee.full_name * 1000000 }}"))
